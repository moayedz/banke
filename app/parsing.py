"""قراءة كشوفات الحساب (Excel / CSV) وتحويلها لمعاملات.

البنوك التركية كل واحد إله شكل، فبدل ما نكتب قارئ لكل بنك، بندوّر على
صف العناوين (Tarih / Açıklama / Tutar / Borç / Alacak ...) وبنفهم الأعمدة منه.
"""

import csv
import hashlib
import io
import re
from datetime import date, datetime

_TR_MAP = str.maketrans({"ı": "i", "ş": "s", "ç": "c", "ğ": "g", "ü": "u", "ö": "o", "â": "a", "î": "i", "û": "u"})


def normalize(text: str) -> str:
    """نص صغير بدون أحرف تركية خاصة، عشان المقارنة تكون سهلة."""
    text = str(text or "").replace("İ", "i").replace("I", "ı").lower()
    text = text.translate(_TR_MAP)
    return re.sub(r"\s+", " ", text).strip()


# ---------- المبالغ والتواريخ ----------

def parse_amount(value) -> float | None:
    """بيفهم '1.234,56' و '-1.234,56 TL' و '1,234.56' و الأرقام من Excel."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip()
    if not s:
        return None
    negative = s.startswith("-") or s.endswith("-") or (s.startswith("(") and s.endswith(")"))
    s = re.sub(r"[^\d.,]", "", s)
    if not s or not re.search(r"\d", s):
        return None
    if "," in s and "." in s:
        # آخر فاصلة/نقطة هي الفاصلة العشرية
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        if re.search(r",\d{1,2}$", s):
            s = s.replace(",", ".")
        else:
            s = s.replace(",", "")
    elif s.count(".") > 1 or re.search(r"\.\d{3}$", s):
        # '1.234' بالتركي = ألف وميتين...
        s = s.replace(".", "")
    try:
        number = float(s)
    except ValueError:
        return None
    return -number if negative else number


_DATE_FORMATS = ["%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%y", "%d/%m/%y", "%Y.%m.%d"]


def parse_date(value) -> str | None:
    """بيرجّع التاريخ بصيغة YYYY-MM-DD."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    s = str(value).strip()
    s = s.split(" ")[0].split("T")[0]  # نشيل الساعة إذا موجودة
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return None


# ---------- قراءة الملف ----------

def read_table(filename: str, content: bytes) -> list[list]:
    """بيرجّع الملف كجدول (قائمة صفوف)."""
    name = filename.lower()
    if name.endswith(".xlsx"):
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        return [list(row) for row in ws.iter_rows(values_only=True)]
    if name.endswith(".xls"):
        import xlrd

        book = xlrd.open_workbook(file_contents=content)
        sheet = book.sheet_by_index(0)
        rows = []
        for r in range(sheet.nrows):
            row = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                else:
                    row.append(cell.value)
            rows.append(row)
        return rows
    if name.endswith((".csv", ".txt")):
        for encoding in ("utf-8-sig", "cp1254", "latin-1"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t|")
        except csv.Error:
            class dialect(csv.excel):
                delimiter = ";" if text.count(";") > text.count(",") else ","
        return [row for row in csv.reader(io.StringIO(text), dialect)]
    raise ValueError("نوع الملف مش مدعوم. استعمل Excel (.xlsx / .xls) أو CSV.")


# ---------- فهم الأعمدة ----------

COLUMN_HINTS = {
    "date": ["islem tarihi", "tarih", "date", "valor"],
    "description": ["aciklama", "islem aciklamasi", "description", "detay", "islem", "aciklamalar"],
    "amount": ["islem tutari", "tutar", "miktar", "amount"],
    "debit": ["borc", "cikan", "harcama"],
    "credit": ["alacak", "giren", "yatan"],
}


def _match_role(header: str) -> str | None:
    h = normalize(header)
    if not h:
        return None
    for role, hints in COLUMN_HINTS.items():
        for hint in hints:
            if h == hint or h.startswith(hint + " ") or h.startswith(hint + "("):
                return role
    if "bakiye" in h:
        return "balance"
    for role, hints in COLUMN_HINTS.items():
        if any(hint in h for hint in hints):
            return role
    return None


def detect_columns(rows: list[list]) -> tuple[int | None, dict[str, int]]:
    """بيدوّر بأول 40 صف على صف العناوين. بيرجّع (رقم الصف، {الدور: رقم العمود})."""
    best: tuple[int | None, dict[str, int]] = (None, {})
    for i, row in enumerate(rows[:40]):
        mapping: dict[str, int] = {}
        for j, cell in enumerate(row):
            role = _match_role(cell) if isinstance(cell, str) else None
            if role and role != "balance" and role not in mapping:
                mapping[role] = j
        has_amount = "amount" in mapping or ("debit" in mapping or "credit" in mapping)
        if "date" in mapping and has_amount and len(mapping) > len(best[1]):
            best = (i, mapping)
    return best


def extract_transactions(rows: list[list], header_row: int, mapping: dict[str, int]) -> list[dict]:
    """بيطلّع المعاملات من الصفوف تحت العناوين، وبيتجاهل الصفوف الفاضية والمجاميع."""

    def cell(row, role):
        idx = mapping.get(role)
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    result = []
    for row in rows[header_row + 1:]:
        tx_date = parse_date(cell(row, "date"))
        if not tx_date:
            continue
        if "amount" in mapping:
            amount = parse_amount(cell(row, "amount"))
        else:
            debit = parse_amount(cell(row, "debit")) or 0
            credit = parse_amount(cell(row, "credit")) or 0
            amount = abs(credit) - abs(debit) if (debit or credit) else None
        if amount is None or amount == 0:
            continue
        description = str(cell(row, "description") or "").strip()
        result.append({"date": tx_date, "description": description, "amount": round(amount, 2)})
    return result


def import_hashes(account_id: int, transactions: list[dict]) -> list[str]:
    """بصمة لكل معاملة عشان إذا استوردت نفس الكشف مرتين ما تتكرّر.

    إذا في معاملتين متطابقتين بنفس الملف (قهوتين بنفس اليوم)، بنزيد رقم تسلسلي.
    """
    seen: dict[str, int] = {}
    hashes = []
    for tx in transactions:
        key = f"{account_id}|{tx['date']}|{tx['amount']:.2f}|{normalize(tx['description'])}"
        seen[key] = seen.get(key, 0) + 1
        hashes.append(hashlib.sha256(f"{key}#{seen[key]}".encode()).hexdigest())
    return hashes


# ---------- التصنيف التلقائي ----------

def categorize(description: str, rules: list[tuple[str, int]]) -> int | None:
    """rules = [(keyword, category_id)]. الكلمة الأطول بتربح."""
    text = normalize(description)
    best = None
    for keyword, category_id in rules:
        kw = normalize(keyword)
        if kw and re.search(r"(?<![a-z0-9])" + re.escape(kw), text):
            if best is None or len(kw) > len(best[0]):
                best = (kw, category_id)
    return best[1] if best else None
